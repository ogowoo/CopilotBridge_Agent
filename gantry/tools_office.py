"""tools_office.py — Office 文档工具（ToolRegistry 的 mixin）：Excel / Word / PowerPoint / PDF 读写。

从 tools.py 拆出（第四批解耦）：方法体原样搬入 mixin，运行时行为不变。
"""
from __future__ import annotations

import asyncio
import base64
import difflib
import hashlib
import inspect
import json
import queue
import re
import shutil
import subprocess
import sys
import threading
import time
from datetime import datetime
from pathlib import Path
from typing import Any, Awaitable, Callable

from . import editmatch

from .tools import ToolError  # noqa: F401


class OfficeMixin:
    def _office_allowed(self):
        if not self.allow_office:
            raise ToolError("Office 工具已在 config.json 中禁用（allow_office=false）")

    @staticmethod
    def _int_arg(value, default: int, name: str) -> int:
        try:
            return int(value)
        except (TypeError, ValueError):
            raise ToolError(f"{name} 必须是整数")

    async def read_excel(self, path: str, sheet: str = "", max_rows: int = 200) -> dict[str, Any]:
        self._office_allowed()
        target = self._path(path)
        if not target.is_file():
            raise ToolError(f"文件不存在: {path}")
        try:
            import openpyxl
        except ImportError:
            raise ToolError("缺少依赖 openpyxl，请运行 pip install openpyxl")
        limit = max(1, min(self._int_arg(max_rows, 200, "max_rows"), 2000))
        wb = openpyxl.load_workbook(target, data_only=True, read_only=True)
        try:
            if sheet and sheet not in wb.sheetnames:
                raise ToolError(f"工作表不存在: {sheet}，可用: {wb.sheetnames}")
            out = {}
            for name in ([sheet] if sheet else wb.sheetnames):
                ws = wb[name]
                rows, total = [], 0
                for r in ws.iter_rows(values_only=True):
                    total += 1
                    if len(rows) < limit:
                        rows.append(["" if c is None else str(c) for c in r])
                out[name] = {"rows": rows, "total_rows": total, "truncated": total > limit}
            return {"path": path, "sheets": out}
        finally:
            wb.close()

    async def write_excel(self, path: str, rows: list, sheet: str = "", start_cell: str = "A1") -> dict[str, Any]:
        self._office_allowed()
        if not isinstance(rows, list) or not rows:
            raise ToolError("rows 必须是非空二维数组")
        try:
            import openpyxl
            from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
        except ImportError:
            raise ToolError("缺少依赖 openpyxl，请运行 pip install openpyxl")
        target = self._path(path)
        self._snap("write_excel", path, target)
        target.parent.mkdir(parents=True, exist_ok=True)
        before = self._hash(target)
        try:
            wb = openpyxl.load_workbook(target) if target.is_file() else openpyxl.Workbook()
        except Exception as exc:
            raise ToolError(f"无法读取 Excel 文件（可能不是有效的 xlsx）: {exc}")
        try:
            if sheet:
                ws = wb[sheet] if sheet in wb.sheetnames else wb.create_sheet(sheet)
            else:
                ws = wb.active
            try:
                col_s, row_s = coordinate_from_string(start_cell)
                col0 = column_index_from_string(col_s)
            except Exception:
                raise ToolError(f"无效的起始单元格: {start_cell}")
            written = 0
            for ri, row in enumerate(rows):
                if not isinstance(row, list):
                    row = [row]
                for ci, value in enumerate(row):
                    ws.cell(row=row_s + ri, column=col0 + ci, value=value)
                    written += 1
            wb.save(target)
            title = ws.title
        finally:
            wb.close()
        return {"path": path, "sheet": title, "cells_written": written,
                "file_event": {"tool": "write_excel", "path": path,
                               "before_hash": before, "after_hash": self._hash(target), "result": "ok"}}

    async def read_docx(self, path: str) -> dict[str, Any]:
        self._office_allowed()
        target = self._path(path)
        if not target.is_file():
            raise ToolError(f"文件不存在: {path}")
        try:
            import docx
        except ImportError:
            raise ToolError("缺少依赖 python-docx，请运行 pip install python-docx")
        doc = docx.Document(target)
        parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text)
        for ti, table in enumerate(doc.tables, 1):
            parts.append(f"[表格 {ti}]")
            for row in table.rows:
                parts.append(" | ".join(cell.text.strip() for cell in row.cells))
        text = "\n".join(parts)
        return {"path": path, "paragraphs": len(doc.paragraphs), "tables": len(doc.tables),
                "content": text[:self.max_read], "truncated": len(text) > self.max_read}

    async def write_docx(self, path: str, paragraphs: list, mode: str = "append") -> dict[str, Any]:
        self._office_allowed()
        if not isinstance(paragraphs, list) or not paragraphs:
            raise ToolError("paragraphs 必须是非空字符串数组")
        if mode not in ("append", "overwrite"):
            raise ToolError("mode 只能是 append 或 overwrite")
        try:
            import docx
        except ImportError:
            raise ToolError("缺少依赖 python-docx，请运行 pip install python-docx")
        target = self._path(path)
        self._snap("write_docx", path, target)
        target.parent.mkdir(parents=True, exist_ok=True)
        before = self._hash(target)
        if mode == "overwrite" or not target.is_file():
            doc = docx.Document()
        else:
            doc = docx.Document(target)
        for text in paragraphs:
            doc.add_paragraph(str(text))
        doc.save(target)
        return {"path": path, "mode": mode, "paragraphs_added": len(paragraphs),
                "file_event": {"tool": "write_docx", "path": path,
                               "before_hash": before, "after_hash": self._hash(target), "result": "ok"}}

    async def read_pptx(self, path: str) -> dict[str, Any]:
        self._office_allowed()
        target = self._path(path)
        if not target.is_file():
            raise ToolError(f"文件不存在: {path}")
        try:
            from pptx import Presentation
        except ImportError:
            raise ToolError("缺少依赖 python-pptx，请运行 pip install python-pptx")
        prs = Presentation(target)
        slides = []
        for i, slide in enumerate(prs.slides, 1):
            texts = [shape.text for shape in slide.shapes if shape.has_text_frame and shape.text.strip()]
            notes = ""
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                notes = slide.notes_slide.notes_text_frame.text.strip()
            slides.append({"slide": i, "text": "\n".join(texts), "notes": notes})
        total_chars = sum(len(s["text"]) + len(s["notes"]) for s in slides)
        return {"path": path, "slides": slides, "slide_count": len(slides),
                "truncated": total_chars > self.max_read}

    async def read_pdf(self, path: str) -> dict[str, Any]:
        self._office_allowed()
        target = self._path(path)
        if not target.is_file():
            raise ToolError(f"文件不存在: {path}")
        try:
            from pypdf import PdfReader
        except ImportError:
            raise ToolError("缺少依赖 pypdf，请运行 pip install pypdf")
        reader = PdfReader(str(target))
        pages = []
        used = 0
        truncated = False
        for i, page in enumerate(reader.pages, 1):
            if used >= self.max_read:
                truncated = True
                break
            text = (page.extract_text() or "").strip()
            used += len(text)
            pages.append({"page": i, "text": text})
        return {"path": path, "page_count": len(reader.pages), "pages_read": len(pages),
                "pages": pages, "truncated": truncated or len(pages) < len(reader.pages)}
